"""Exercise only a uniquely named, isolated Compose mock deployment.

Requires Docker Engine, backend Python dependencies, built frontend test
dependencies and a Playwright browser. Test volumes are retained, never removed.
Use --state to resume this script's own alem-verify-* project after a failed run.
"""

import argparse
import hashlib
import io
import json
import os
import re
import secrets
import shutil
import subprocess
import time
import wave
from contextlib import contextmanager
from pathlib import Path
from uuid import uuid4

import httpx
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
EDIT = "Проверенное поручение из браузера"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--state", type=Path)
    parser.add_argument("--no-build", action="store_true")
    args = parser.parse_args()
    if args.state:
        state = json.loads(args.state.read_text(encoding="utf-8"))
    else:
        suffix = uuid4().hex[:8]
        artifacts = ROOT / ".artifacts" / ("deployment-" + suffix)
        artifacts.mkdir(parents=True)
        state = {"project": "alem-verify-" + suffix, "artifacts": str(artifacts), "env_file": str(artifacts / ".env"), "web_port": 8088}
        values = {"APP_ENV": "test", "COOKIE_SECURE": "false", "FRONTEND_ORIGIN": "http://127.0.0.1:8088", "WEB_PORT": "8088", "POSTGRES_PASSWORD": secrets.token_hex(24), "AI_INTERNAL_TOKEN": secrets.token_hex(32), "AI_EXPECTED_MODE": "mock", "TEST_EMAIL": "browser@example.test", "TEST_PASSWORD": secrets.token_urlsafe(24)}
        Path(state["env_file"]).write_text("\n".join(f"{k}={v}" for k, v in values.items()) + "\n", encoding="utf-8")
        (artifacts / "state.json").write_text(json.dumps(state, indent=2), encoding="utf-8")
    assert re.fullmatch(r"alem-verify-[0-9a-f]{8}", state["project"]), "Refusing to operate on a non-test project"
    artifacts = Path(state["artifacts"]).resolve()
    assert artifacts.is_relative_to(ROOT / ".artifacts")
    values = dict(line.split("=", 1) for line in Path(state["env_file"]).read_text(encoding="utf-8").splitlines() if line and not line.startswith("#"))
    env = dict(os.environ, COMPOSE_BAKE="false", COMPOSE_PARALLEL_LIMIT="1", TEST_PASSWORD=values["TEST_PASSWORD"], TEST_EMAIL=values["TEST_EMAIL"], TEST_URL=values["FRONTEND_ORIGIN"], TEST_DOWNLOAD_PATH=str(artifacts / "container-mock.docx"), TEST_RELOAD_PROCESSING="false")
    docker = shutil.which("docker")
    if not docker:
        raise SystemExit("BLOCKED: Docker CLI required")
    prefix = [docker, "compose", "-p", state["project"], "--env-file", state["env_file"], "-f", str(ROOT / "compose.yml"), "--profile", "mock"]
    restore_prefix = None
    report = {"project": state["project"], "status": "FAILED", "mode": "HTTP_AI_MOCK", "checks": [], "volumes_retained": True}

    def command(arguments, timeout=300, child_env=env):
        result = subprocess.run(arguments, cwd=ROOT, env=child_env, text=True, encoding="utf-8", errors="replace", stdout=subprocess.PIPE, stderr=subprocess.STDOUT, timeout=timeout)
        output = result.stdout
        for key in ("POSTGRES_PASSWORD", "AI_INTERNAL_TOKEN", "TEST_PASSWORD"):
            output = output.replace(values[key], "[REDACTED]")
        with (artifacts / "container-checks.log").open("a", encoding="utf-8") as log:
            log.write(output + "\n")
        if result.returncode:
            raise RuntimeError(f"Command failed ({result.returncode}); see container-checks.log")
        return output

    def compose(*arguments, **kwargs):
        return command([*prefix, *arguments], **kwargs)

    def passed(name, **details):
        report["checks"].append({"name": name, "status": "PASSED", **details})
        print("PASSED: " + name, flush=True)

    def ready(base):
        deadline = time.monotonic() + 180
        while time.monotonic() < deadline:
            try:
                if httpx.get(base + "/api/ready", timeout=6, trust_env=False).status_code == 200:
                    return
            except httpx.RequestError:
                pass
            time.sleep(2)
        raise RuntimeError("Readiness deadline exceeded")

    @contextmanager
    def login(base):
        with httpx.Client(base_url=base + "/api", timeout=15, trust_env=False, headers={"Origin": base}) as client:
            response = client.post("/auth/login", json={"email": values["TEST_EMAIL"], "password": values["TEST_PASSWORD"]})
            response.raise_for_status()
            client.headers["X-CSRF-Token"] = response.json()["csrf_token"]
            yield client

    def validate_saved(client, meeting_id, audio_hash):
        detail = client.get("/meetings/" + meeting_id).json()
        assert any(item["text"] == EDIT for item in detail["action_items"])
        audio = client.get(f"/meetings/{meeting_id}/audio")
        audio.raise_for_status()
        assert hashlib.sha256(audio.content).hexdigest() == audio_hash
        partial = client.get(f"/meetings/{meeting_id}/audio", headers={"Range": "bytes=0-15"})
        assert partial.status_code == 206 and len(partial.content) == 16

    def check_docx(content):
        document = Document(io.BytesIO(content))
        assert any("Демонстрационный результат" in p.text for p in document.paragraphs)
        assert EDIT in " ".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)

    try:
        info = command([docker, "info", "--format", "{{json .ServerVersion}}|{{json .OSType}}|{{json .Architecture}}|{{json .MemTotal}}|{{json .NCPU}}"])
        version, os_type, architecture, memory, cpus = map(json.loads, info.strip().split("|"))
        report["platform"] = {"ServerVersion": version, "OSType": os_type, "Architecture": architecture, "MemTotal": memory, "NCPU": cpus}
        compose("config", "--quiet")
        if not args.no_build:
            # Compose v5 may delegate its multi-target build to Bake. Build each
            # image sequentially here to keep peak memory bounded on Windows.
            command([docker, "build", "-f", "backend/Dockerfile", "-t", state["project"] + "-api:latest", "."], timeout=900)
            for service in ("worker", "migrate"):
                command([docker, "tag", state["project"] + "-api:latest", state["project"] + f"-{service}:latest"])
            command([docker, "build", "-f", "frontend/Dockerfile", "-t", state["project"] + "-web:latest", "."], timeout=900)
            command([docker, "build", "-f", "ai-service/Dockerfile", "--target", "base", "-t", state["project"] + "-ai-mock:latest", "."], timeout=900)
        compose("up", "-d", "--no-build", "--wait", "--wait-timeout", "180")
        base = values["FRONTEND_ORIGIN"]
        ready(base)
        passed("build/start/migrations/readiness", build_executed=not args.no_build)
        existing = compose("exec", "-T", "db", "psql", "-U", "alem", "-d", "alem", "-Atc", "SELECT count(*) FROM users WHERE email = 'browser@example.test'").strip()
        if existing == "0":
            compose("exec", "-T", "-e", "TEST_PASSWORD", "api", "python", "-m", "app.manage", "create-user", "--email", values["TEST_EMAIL"], "--name", "Container test", "--password-env", "TEST_PASSWORD")
        command([shutil.which("node"), str(ROOT / "frontend/tests/connected.mjs")], timeout=150)
        check_docx(Path(env["TEST_DOWNLOAD_PATH"]).read_bytes())
        passed("browser upload/edit/reload/confirm/DOCX/relogin through Nginx")
        with login(base) as client:
            history = client.get("/meetings").json()["items"]
            meeting_id = next(item["id"] for item in history if any(task["text"] == EDIT for task in client.get("/meetings/" + item["id"]).json()["action_items"]))
            audio_hash = hashlib.sha256(client.get(f"/meetings/{meeting_id}/audio").content).hexdigest()
            compose("restart", "-t", "15", "db", "api", "worker", "ai-mock", "web")
            ready(base)
            validate_saved(client, meeting_id, audio_hash)
            passed("database/audio/session/edit survive container restart")

            # Pause the actual AI container to hold an in-flight request. No API stub.
            compose("pause", "ai-mock")
            try:
                audio = io.BytesIO()
                with wave.open(audio, "wb") as wav:
                    wav.setparams((1, 2, 16000, 0, "NONE", "not compressed"))
                    wav.writeframes(bytes(32000))
                response = client.post("/meetings", data={"title": "Worker crash recovery", "started_at": "2026-09-25T10:00:00+05:00", "timezone": "Asia/Almaty", "participants_json": json.dumps([{"name": "Test"}])}, files={"file": ("recovery.wav", audio.getvalue(), "audio/wav")})
                response.raise_for_status()
                recovery_id = response.json()["id"]
                deadline = time.monotonic() + 30
                while client.get("/meetings/" + recovery_id).json()["status"] != "running":
                    assert time.monotonic() < deadline, "Worker did not claim recovery job"
                    time.sleep(1)
                compose("kill", "-s", "SIGKILL", "worker")
            finally:
                compose("unpause", "ai-mock")
            compose("up", "-d", "--no-build", "worker")
            deadline = time.monotonic() + 250
            print("Waiting for the configured 180-second worker lease to expire", flush=True)
            while True:
                detail = client.get("/meetings/" + recovery_id).json()
                assert detail["status"] != "failed", "Recovered job failed"
                if detail["status"] == "completed":
                    break
                assert time.monotonic() < deadline, "Lease recovery deadline exceeded"
                time.sleep(3)
            counts = compose("exec", "-T", "db", "psql", "-U", "alem", "-d", "alem", "-Atc", f"SELECT count(*) FROM jobs WHERE meeting_id = '{recovery_id}'").strip()
            assert counts == "1"
            validate_saved(client, meeting_id, audio_hash)
            passed("worker SIGKILL/180-second lease recovery without duplicate job or overwritten edit")

        compose("stop", "-t", "20", "web", "api", "worker")
        compose("exec", "-T", "db", "pg_dump", "-U", "alem", "-d", "alem", "-Fc", "-f", "/tmp/verification.dump")
        compose("cp", "db:/tmp/verification.dump", str(artifacts / "db.dump"))
        backup_audio = artifacts / ("audio-backup-" + uuid4().hex[:8])
        compose("cp", "api:/data/audio/.", str(backup_audio))
        restore_project = "alem-restore-" + uuid4().hex[:8]
        override = artifacts / "restore-images.json"
        override.write_text(json.dumps({"services": {name: {"image": f'{state["project"]}-{name}:latest'} for name in ("api", "worker", "migrate", "web", "ai-mock")}}), encoding="utf-8")
        restore_env = dict(env, WEB_PORT="8089", FRONTEND_ORIGIN="http://127.0.0.1:8089")
        restore_prefix = [docker, "compose", "-p", restore_project, "--env-file", state["env_file"], "-f", str(ROOT / "compose.yml"), "-f", str(override), "--profile", "mock"]

        def restore(*arguments):
            return command([*restore_prefix, *arguments], child_env=restore_env)

        report["restore_project"] = restore_project
        restore("up", "-d", "--no-build", "--wait", "db")
        restore("cp", str(artifacts / "db.dump"), "db:/tmp/verification.dump")
        restore("exec", "-T", "db", "pg_restore", "-U", "alem", "-d", "alem", "/tmp/verification.dump")
        restore("create", "--no-build", "api", "worker")
        restore("cp", str(backup_audio) + os.sep + ".", "api:/data/audio")
        restore("run", "--rm", "--no-deps", "--user", "root", "api", "chown", "-R", "appuser:appuser", "/data/audio")
        restore("up", "-d", "--no-build", "--wait", "--wait-timeout", "180")
        ready(restore_env["FRONTEND_ORIGIN"])
        with login(restore_env["FRONTEND_ORIGIN"]) as client:
            validate_saved(client, meeting_id, audio_hash)
            client.post(f"/meetings/{meeting_id}/confirm").raise_for_status()
            response = client.get(f"/meetings/{meeting_id}/export?format=docx")
            response.raise_for_status()
            check_docx(response.content)
        passed("isolated second-project pg_restore/audio hash/owner login/Range/DOCX")
        report["status"] = "PASSED"
    finally:
        # Only this script's project names; keep all volumes and backup files.
        for current in (restore_prefix, prefix):
            if current:
                try:
                    command([*current, "stop", "-t", "20"], timeout=90)
                except Exception:
                    report.setdefault("cleanup_errors", []).append("Could not stop " + current[3])
        (artifacts / "container-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f'{report["status"]}: {artifacts / "container-report.json"}', flush=True)


if __name__ == "__main__":
    main()
