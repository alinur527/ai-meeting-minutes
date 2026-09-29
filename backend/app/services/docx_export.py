from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


@dataclass(frozen=True)
class ExportTask:
    text: str
    assignee: str | None
    due_date: date | None
    deadline_text: str | None = None
    review_reason: str | None = None


@dataclass(frozen=True)
class ExportSegment:
    speaker: str
    start_ms: int
    end_ms: int
    text: str


@dataclass(frozen=True)
class MeetingExportData:
    title: str
    started_at: datetime
    timezone: str
    participants: tuple[str, ...]
    summary: str | None
    tasks: tuple[ExportTask, ...]
    segments: tuple[ExportSegment, ...]
    processing_mode: str | None = None


def _set_cell_shading(cell, fill: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        cell_properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_borders(cell, color: str = "D9D9D9") -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    borders = cell_properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        cell_properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def _set_cell_margins(cell, margin_twips: int = 110) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)
    for side in ("top", "start", "bottom", "end"):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(margin_twips))
        element.set(qn("w:type"), "dxa")


def _format_timestamp(milliseconds: int) -> str:
    total_seconds, millis = divmod(milliseconds, 1000)
    minutes, seconds = divmod(total_seconds, 60)
    hours, minutes = divmod(minutes, 60)
    if hours:
        return f"{hours:02d}:{minutes:02d}:{seconds:02d}.{millis:03d}"
    return f"{minutes:02d}:{seconds:02d}.{millis:03d}"


def _set_font(run, *, size: float = 11, bold: bool = False, color=None) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def _configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = document.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_properties = title._element.get_or_add_pPr()
    title_borders = title_properties.find(qn("w:pBdr"))
    if title_borders is not None:
        title_properties.remove(title_borders)

    for style_name in ("Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.color.rgb = RGBColor(0, 0, 0)


def _style_table(table, widths: tuple[float, ...]) -> None:
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_borders(cell)
            _set_cell_margins(cell)
            if row_index == 0:
                _set_cell_shading(cell, "1F4E78")
                for run in cell.paragraphs[0].runs:
                    _set_font(run, bold=True, color=(255, 255, 255))
            elif row_index % 2 == 0:
                _set_cell_shading(cell, "EAF2F8")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    if row_index != 0:
                        _set_font(run)


def build_meeting_docx(data: MeetingExportData) -> BytesIO:
    """Build an in-memory protocol document from a detached database snapshot."""

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    _configure_styles(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run(data.title)
    if data.processing_mode != "real":
        document.add_paragraph("Демонстрационный результат: запись не распознавалась" if data.processing_mode in {"ai_mock", "backend_mock"} else "Режим обработки исходной записи не подтверждён")

    local_started_at = data.started_at.astimezone(ZoneInfo(data.timezone))
    metadata = document.add_paragraph()
    run = metadata.add_run("Дата и время: ")
    _set_font(run, bold=True)
    _set_font(metadata.add_run(local_started_at.isoformat(timespec="minutes")))
    metadata.add_run("\n")
    run = metadata.add_run("Часовой пояс: ")
    _set_font(run, bold=True)
    _set_font(metadata.add_run(data.timezone))

    document.add_heading("Участники", level=1)
    if data.participants:
        for participant in data.participants:
            paragraph = document.add_paragraph(style="List Bullet")
            _set_font(paragraph.add_run(participant))
    else:
        document.add_paragraph("Участники не указаны")

    document.add_heading("Краткое содержание", level=1)
    document.add_paragraph(data.summary or "Краткое содержание не указано")

    document.add_heading("Поручения", level=1)
    task_table = document.add_table(rows=1, cols=3)
    task_table.rows[0].cells[0].text = "Поручение"
    task_table.rows[0].cells[1].text = "Ответственный"
    task_table.rows[0].cells[2].text = "Срок"
    if data.tasks:
        for task in data.tasks:
            cells = task_table.add_row().cells
            cells[0].text = task.text
            if task.review_reason:
                cells[0].text += "\nТребует проверки: " + task.review_reason
            cells[1].text = task.assignee or "Не указан"
            cells[2].text = task.due_date.isoformat() if task.due_date else "Не указан"
            if task.deadline_text:
                cells[2].text += "\nИсходный срок: " + task.deadline_text
    else:
        cells = task_table.add_row().cells
        cells[0].text = "Поручения отсутствуют"
        cells[1].text = "Не указан"
        cells[2].text = "Не указан"
    _style_table(task_table, (3.5, 1.8, 1.6))

    document.add_heading("Расшифровка", level=1)
    transcript_table = document.add_table(rows=1, cols=3)
    transcript_table.rows[0].cells[0].text = "Спикер"
    transcript_table.rows[0].cells[1].text = "Время"
    transcript_table.rows[0].cells[2].text = "Текст"
    if data.segments:
        for segment in data.segments:
            cells = transcript_table.add_row().cells
            cells[0].text = segment.speaker
            cells[1].text = (
                f"{_format_timestamp(segment.start_ms)} - "
                f"{_format_timestamp(segment.end_ms)}"
            )
            cells[2].text = segment.text
    else:
        cells = transcript_table.add_row().cells
        cells[0].text = "Не указан"
        cells[1].text = "Не указан"
        cells[2].text = "Расшифровка отсутствует"
    _style_table(transcript_table, (1.55, 1.65, 3.7))

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
