"""Full isolated stack: real PostgreSQL/API/worker/static frontend + explicit AI mock."""

import hashlib
import os
import secrets
import shutil
import socket
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from uuid import uuid4

import httpx
import psycopg
from docx import Document
from psycopg import sql
from sqlalchemy.engine import make_url

ROOT = Path(__file__).resolve().parents[2]


def port():
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


def wait_ready(url, process, headers=None):
    with httpx.Client(trust_env=False, timeout=2) as client:
        for _ in range(100):
            if process.poll() is not None:
                raise RuntimeError("Service exited; inspect .artifacts/e2e logs")
            try:
                if client.get(url, headers=headers).status_code == 200:
                    return
            except httpx.RequestError:
                pass
            time.sleep(0.2)
    raise RuntimeError(f"Service did not become ready at {url}")


def main():
    ai_python = os.environ["AI_PYTHON"]
    pg_bin = os.environ.get("PG_BIN")
    if not pg_bin or any(not shutil.which(command, path=pg_bin) for command in ("pg_dump", "pg_restore")):
        raise SystemExit("BLOCKED: PG_BIN must contain pg_dump and pg_restore; backup verification is mandatory")
    admin_url = make_url(os.environ["TEST_POSTGRES_ADMIN_URL"])
    name = "alem_test_" + uuid4().hex
    dsn = admin_url.set(drivername="postgresql").render_as_string(hide_password=False)
    db_url = admin_url.set(database=name).render_as_string(hide_password=False)
    api_port, ai_port, web_port = port(), port(), port()
    origin = f"http://127.0.0.1:{web_port}"
    artifacts = Path(os.environ.get("E2E_ARTIFACT_DIR", ROOT / ".artifacts" / "e2e"))
    artifacts.mkdir(parents=True, exist_ok=True)
    processes, logs = [], []
    with (
        tempfile.TemporaryDirectory(prefix="alem-e2e-") as temporary,
        psycopg.connect(dsn, autocommit=True) as admin,
    ):
        storage = Path(temporary) / "audio"
        storage.mkdir()
        env = dict(
            os.environ,
            APP_ENV="test",
            DATABASE_URL=db_url,
            COOKIE_SECURE="false",
            STORAGE_DIR=str(storage),
            FRONTEND_ORIGIN=origin,
            AI_MODE="http",
            AI_EXPECTED_MODE="mock",
            AI_BASE_URL=f"http://127.0.0.1:{ai_port}",
            AI_INTERNAL_TOKEN=secrets.token_hex(32),
            TEST_EMAIL="browser@example.test",
            TEST_PASSWORD=secrets.token_urlsafe(20),
            TEST_URL=origin,
            BACKEND_URL=f"http://127.0.0.1:{api_port}",
            VITE_USE_MOCK="false",
            TEST_DOWNLOAD_PATH=str(artifacts / "protocol.docx"),
            TEST_RELOAD_PROCESSING="true",
        )

        def start(label, args, cwd, child_env):
            log = (artifacts / f"{label}.log").open("w", encoding="utf-8")
            logs.append(log)
            process = subprocess.Popen(
                args, cwd=cwd, env=child_env, stdout=log, stderr=log
            )
            processes.append(process)
            return process

        admin.execute(
            sql.SQL("CREATE DATABASE {} ENCODING 'UTF8' TEMPLATE template0").format(
                sql.Identifier(name)
            )
        )
        try:
            subprocess.run(
                [sys.executable, "-m", "alembic", "upgrade", "head"],
                cwd=ROOT / "backend",
                env=env,
                check=True,
            )
            subprocess.run(
                [
                    sys.executable,
                    "-m",
                    "app.manage",
                    "create-user",
                    "--email",
                    env["TEST_EMAIL"],
                    "--name",
                    "Browser test",
                    "--password-env",
                    "TEST_PASSWORD",
                ],
                cwd=ROOT / "backend",
                env=env,
                check=True,
            )
            # Test-only latency makes reload-during-processing deterministic.
            wrapper = Path(temporary) / "e2e_ai.py"
            wrapper.write_text(
                "import time\nfrom app.main import app\noriginal = app.state.pipeline.process\ndef delayed(path, meta):\n    time.sleep(3)\n    return original(path, meta)\napp.state.pipeline.process = delayed\n",
                encoding="utf-8",
            )
            ai_env = dict(
                env,
                AI_MODE="mock",
                INTERNAL_TOKEN=env["AI_INTERNAL_TOKEN"],
                PYTHONPATH=os.pathsep.join([temporary, str(ROOT / "ai-service")]),
            )
            ai = start(
                "ai",
                [
                    ai_python,
                    "-m",
                    "uvicorn",
                    "e2e_ai:app",
                    "--host",
                    "127.0.0.1",
                    "--port",
                    str(ai_port),
                ],
                ROOT / "ai-service",
                ai_env,
            )
            wait_ready(
                env["AI_BASE_URL"] + "/ready",
                ai,
                {"X-Internal-Token": env["AI_INTERNAL_TOKEN"]},
            )
            api = start(
                "api",
                [
                    sys.executable,
                    "-m",
                    "uvicorn",
                    "app.main:app",
                    "--host",
                    "127.0.0.1",
                    "--port",
                    str(api_port),
                ],
                ROOT / "backend",
                env,
            )
            start("worker", [sys.executable, "-m", "app.worker"], ROOT / "backend", env)
            wait_ready(env["BACKEND_URL"] + "/ready", api)
            web = start(
                "web",
                [
                    "node",
                    "node_modules/vite/bin/vite.js",
                    "preview",
                    "--host",
                    "127.0.0.1",
                    "--port",
                    str(web_port),
                ],
                ROOT / "frontend",
                env,
            )
            wait_ready(origin, web)
            subprocess.run(
                ["node", "tests/connected.mjs"],
                cwd=ROOT / "frontend",
                env=env,
                check=True,
                timeout=120,
            )
            document = Document(env["TEST_DOWNLOAD_PATH"])
            contents = " ".join(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
            assert "Проверенное поручение из браузера" in contents
            assert any(
                "Демонстрационный результат" in paragraph.text
                for paragraph in document.paragraphs
            )
            print("PASSED: downloaded DOCX contains the persisted browser edit")

            # Quiesce writers before taking a coordinated database/audio snapshot.
            for process in processes:
                process.terminate()
            for process in processes:
                process.wait(timeout=10)
            if pg_bin:
                restored = name + "_restore"
                pg_env = dict(os.environ, PGPASSWORD=admin_url.password or "")
                args = [
                    "-h",
                    admin_url.host,
                    "-p",
                    str(admin_url.port or 5432),
                    "-U",
                    admin_url.username,
                ]
                dump = Path(temporary) / "db.dump"
                subprocess.run(
                    [
                        str(Path(pg_bin) / "pg_dump"),
                        *args,
                        "-d",
                        name,
                        "-Fc",
                        "-f",
                        str(dump),
                    ],
                    env=pg_env,
                    check=True,
                )
                copied = Path(temporary) / "restored-audio"
                shutil.copytree(storage, copied)
                admin.execute(
                    sql.SQL("CREATE DATABASE {}").format(sql.Identifier(restored))
                )
                try:
                    subprocess.run(
                        [
                            str(Path(pg_bin) / "pg_restore"),
                            *args,
                            "-d",
                            restored,
                            str(dump),
                        ],
                        env=pg_env,
                        check=True,
                    )
                    with psycopg.connect(
                        admin_url.set(
                            drivername="postgresql", database=restored
                        ).render_as_string(hide_password=False)
                    ) as restored_db:
                        paths = restored_db.execute(
                            "SELECT audio_path FROM meetings"
                        ).fetchall()
                        assert paths
                        for (audio_path,) in paths:
                            original = Path(audio_path)
                            assert (
                                hashlib.sha256(original.read_bytes()).digest()
                                == hashlib.sha256(
                                    (copied / original.name).read_bytes()
                                ).digest()
                            )
                        assert (
                            restored_db.execute(
                                "SELECT count(*) FROM action_items WHERE text = %s",
                                ("Проверенное поручение из браузера",),
                            ).fetchone()[0]
                            == 1
                        )
                    print(
                        "PASSED: isolated pg_dump/pg_restore and matching audio backup"
                    )
                finally:
                    admin.execute(
                        sql.SQL("DROP DATABASE {} WITH (FORCE)").format(
                            sql.Identifier(restored)
                        )
                    )
            else:
                print(
                    "BLOCKED: backup restore check needs PG_BIN (PostgreSQL 17+ clients)"
                )
        finally:
            for process in processes:
                if process.poll() is None:
                    process.terminate()
                    try:
                        process.wait(timeout=10)
                    except subprocess.TimeoutExpired:
                        process.kill()
                        process.wait(timeout=5)
            for log in logs:
                log.close()
            admin.execute(
                sql.SQL("DROP DATABASE {} WITH (FORCE)").format(sql.Identifier(name))
            )


if __name__ == "__main__":
    main()
